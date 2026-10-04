import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.properties import Outline

data = json.load(open('data.json'))
scenes = json.load(open('scenes.json'))

F = 'Arial'
def font(**k):
    k.setdefault('name', F); k.setdefault('size', 10); return Font(**k)
BLUE = font(color='0000FF')
BLACK = font()
BOLD = font(bold=True)
H1 = font(bold=True, size=16)
H2 = font(bold=True, size=12)
GREY = font(color='666666', italic=True)
GREEN = font(color='008000')
HEAD_FILL = PatternFill('solid', fgColor='DCE6F1')
INPUT_FILL = PatternFill('solid', fgColor='FFFFCC')
MADE_FILL = PatternFill('solid', fgColor='FBE2C8')   # orange-ish (bid made)
SET_FILL = PatternFill('solid', fgColor='CFEBE8')    # teal-ish (bid set)
PLAN_FILL = PatternFill('solid', fgColor='EADCF3')   # purple-ish (seat 2's plan node)
thin = Side(style='thin', color='BBBBBB')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

# ---------------------------------------------------------------- helpers
def header(ws, row, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=h); c.font = BOLD; c.fill = HEAD_FILL; c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def put(ws, ref, v, f=BLACK, fill=None, al=None, nf=None):
    c = ws[ref]; c.value = v; c.font = f
    if fill: c.fill = fill
    if al: c.alignment = al
    if nf: c.number_format = nf
    return c

# ---------------------------------------------------------------- 1. Read me
rm = wb.active; rm.title = 'Read me'
widths(rm, [3, 100])
put(rm, 'B2', 'Walt in Excel — All the Way Down', H1)
put(rm, 'B3', 'One decision by a Texas 42 player, laid out so every number can be clicked on.', GREY)
lines = [
 ('What Walt is', H2),
 ("Walt is a Texas 42 player. What it does fits in one sentence: it guesses the hidden hands a few times, "
  "plays each guess out to the end assuming everyone else plays at random, and picks the tile that comes out best "
  "across the guesses. That's it. No lookup tables, no learned weights — just counting.", BLACK),
 ('', BLACK),
 ('What this workbook shows', H2),
 ("Trick 5 of a hand. Seat 1 (walt) led the 6-4 and seat 2 has to answer. Seat 2 is defending: the bidders have 18 points, "
  "the defenders 1, and the bid is 30. Seat 2 holds the 6-3, the 6-0 and the 5-5. It must follow sixes, so the choice is 6-3 or 6-0.", BLACK),
 ("Seat 2 can't see the other eight tiles, so it guessed eight ways they might be dealt (the 'Worlds' sheet). "
  "Then for each candidate tile it played every world out to the end, one position at a time. Those positions are the rows on the "
  "'Tree 6-3' and 'Tree 6-0' sheets: 151 of them. 'Decision' adds up the answer. 'Walkthrough' is the guided tour.", BLACK),
 ('', BLACK),
 ('The one thing we ever compute', H2),
 ("Every row in the tree answers the same question: in how many of this row's worlds does the bid get made? That count is the row's Value.", BLACK),
 ("There are only three kinds of row, and each has one formula:", BLACK),
 ("   •  dice   — someone else is to play. Each world rolls a die over that seat's legal tiles. Value = SUM of the children (the roll split the worlds into separate piles).", BLACK),
 ("   •  plan   — seat 2's own turn. No die: it tries every legal tile, playing the same tile in every world at once because it can't tell them apart. Value = MIN of the children (seat 2 is defending; it would be MAX if it were bidding).", BLACK),
 ("   •  leaf   — the hand is decided: bidders reached 30 or defenders reached 13. Value = number of worlds if the bid was made, else 0.", BLACK),
 ("Click any Value cell and read its formula. Follow it down and it ends at leaves; follow it up and it ends at the Decision. All the way down.", BLACK),
 ('', BLACK),
 ('Legend', H2),
 ('Blue text = data from Walt\'s trace (the deal guesses and the dice rolls). Black text = a formula. Green text = a link to another sheet.', BLACK),
 ('Orange fill = bid made. Teal fill = bid set. Purple fill = seat 2\'s own plan nodes. Yellow fill = cells you can change (on the "Two rules" sheet).', BLACK),
 ('The tree sheets are outlined: use the +/− buttons at the left edge (or Data ▸ Group ▸ Show Detail) to collapse a subtree and watch its count flow up.', BLACK),
 ('', BLACK),
 ('The flat version', H2),
 ("'Tape', 'Line 1'–'Line 6' and 'Flat Walt' do the same decision with no recursion at all: the dice are pre-rolled onto a tape, each line plays one ordering of seat 2's tiles to the end, and 'Flat Walt' folds the lines back up. With the tape pinned to the trace it reproduces the trees exactly; set Tape!B3 to \"live\" and press F9 for fresh dice.", BLACK),
 ('', BLACK),
 ('Checked: every rolled tile is legal in its world, every plan node tries exactly its legal tiles, and every Value recomputes from the rows beneath it.', GREY),
 ('Source: walt trace rendered in the "All the Way Down" explainer (part two of "One Mind, Eight Worlds"), Oct 2026. Built for Uncle, former Excel pro, from Jason.', GREY),
]
r = 5
for txt, f in lines:
    c = put(rm, f'B{r}', txt, f, al=WRAP); r += 1
rm.sheet_view.showGridLines = False

# ---------------------------------------------------------------- 2. Tiles
tl = wb.create_sheet('Tiles')
widths(tl, [8, 6, 6, 14, 8, 8, 60])
put(tl, 'A1', 'The 28 dominoes, with fives as trump', H2)
put(tl, 'A2', 'Led, a tile is trump if either end is a 5, otherwise its high end names the suit. Following, a non-trump tile belongs to BOTH its ends: the 3-1 follows threes and ones. Rank in a suit = the other end; a double is highest. Count = point value.', GREY)
header(tl, 4, ['Tile', 'High', 'Low', 'Suit when led', 'Rank in that suit', 'Count', 'Note'])
tiles = [(a, b) for a in range(7) for b in range(a + 1)]
tiles.sort(key=lambda t: (-t[0], -t[1]))
TILE_ROWS = {}
r = 5
for a, b in tiles:
    name = f'{a}-{b}'; TILE_ROWS[name] = r
    put(tl, f'A{r}', name, BLUE)
    put(tl, f'B{r}', f'=VALUE(LEFT(A{r},1))')
    put(tl, f'C{r}', f'=VALUE(RIGHT(A{r},1))')
    put(tl, f'D{r}', f'=IF(OR(B{r}=5,C{r}=5),"fives (trump)",B{r})')
    # rank: double = 7, otherwise the pip that is NOT the suit pip
    put(tl, f'E{r}', f'=IF(B{r}=C{r},7,IF(OR(B{r}=5,C{r}=5),B{r}+C{r}-5,C{r}))')
    put(tl, f'F{r}', f'=IF(B{r}+C{r}=5,5,IF(B{r}+C{r}=10,10,0))')
    r += 1
TILE_LAST = r - 1
put(tl, f'A{r+1}', 'Check: count points total', GREY); put(tl, f'F{r+1}', f'=SUM(F5:F{TILE_LAST})')
put(tl, f'G{r+1}', '35 count points + 7 tricks = 42. Bidders make at 30; a defender set is 13 (42 − 30 + 1).', GREY)
tl.freeze_panes = 'A5'

# ---------------------------------------------------------------- 3. Table (position)
tb = wb.create_sheet('Table')
widths(tb, [26, 12, 12, 12, 12, 14, 12, 40])
put(tb, 'A1', 'The table at trick 5', H2)
put(tb, 'A3', 'Trump', BOLD); put(tb, 'B3', 'fives', BLUE)
put(tb, 'A4', 'Bid', BOLD); put(tb, 'B4', 30, BLUE)
put(tb, 'A5', 'Bidders (seats 1 & 3) have', BOLD); put(tb, 'B5', data['t1'], BLUE)
put(tb, 'A6', 'Defenders (seats 0 & 2) have', BOLD); put(tb, 'B6', data['t0'], BLUE)
put(tb, 'A7', 'Defenders set the bid at', BOLD); put(tb, 'B7', '=42-B4+1')
put(tb, 'A8', 'Seat 2 is', BOLD); put(tb, 'B8', 'defending', BLUE)
put(tb, 'C8', '← the only switch in the tree: plan nodes take MIN when defending, MAX when bidding', GREY)
put(tb, 'A10', 'Seat 2\'s hand', BOLD)
for i, t in enumerate(['6-3', '6-0', '5-5']): put(tb, f'{get_column_letter(2+i)}10', t, BLUE)
put(tb, 'A11', 'Led this trick (seat 1, walt)', BOLD); put(tb, 'B11', '6-4', BLUE)
put(tb, 'A12', 'Led suit', BOLD); put(tb, 'B12', f'=INDEX(Tiles!$D$5:$D${TILE_LAST},MATCH(B11,Tiles!$A$5:$A${TILE_LAST},0))', GREEN)
put(tb, 'A13', 'Seat 2 must follow sixes, so its legal tiles are', BOLD)
for i in range(3):
    col = get_column_letter(2 + i)
    put(tb, f'{col}13', f'=IF(INDEX(Tiles!$D$5:$D${TILE_LAST},MATCH({col}10,Tiles!$A$5:$A${TILE_LAST},0))=$B$12,{col}10,"—")', GREEN)
put(tb, 'A15', 'Already played (tricks 1–4)', BOLD)
header(tb, 16, ['Trick', 'Lead', '2nd', '3rd', '4th', 'Winning tile', 'Points', 'Note'])
played = data['played']
for k in range(4):
    rr = 17 + k
    put(tb, f'A{rr}', k + 1, BLUE)
    for j in range(4): put(tb, f'{get_column_letter(2+j)}{rr}', played[4*k + j], BLUE)
    # winning tile: trump beats everything; otherwise a tile that contains the led pip, ranked by its other pip (double highest)
    tiles_rng = f'Tiles!$A$5:$A${TILE_LAST}'
    def hi_of(ref): return f'INDEX(Tiles!$B$5:$B${TILE_LAST},MATCH({ref},{tiles_rng},0))'
    def lo_of(ref): return f'INDEX(Tiles!$C$5:$C${TILE_LAST},MATCH({ref},{tiles_rng},0))'
    def cnt_of(ref): return f'INDEX(Tiles!$F$5:$F${TILE_LAST},MATCH({ref},{tiles_rng},0))'
    ledpip = hi_of(f'B{rr}'); ledtr = f'OR({hi_of(f"B{rr}")}=5,{lo_of(f"B{rr}")}=5)'
    strengths = []
    for j in range(4):
        ref = f'{get_column_letter(2+j)}{rr}'; h = hi_of(ref); l = lo_of(ref)
        strengths.append(f'IF(OR({h}=5,{l}=5),100+IF({h}={l},7,{h}+{l}-5),IF({ledtr},0,IF(OR({h}={ledpip},{l}={ledpip}),IF({h}={l},7,{h}+{l}-{ledpip}),0)))')
    s = ','.join(strengths)
    put(tb, f'F{rr}', f'=INDEX(B{rr}:E{rr},MATCH(MAX({s}),CHOOSE({{1,2,3,4}},{s}),0))', GREEN)
    put(tb, f'G{rr}', '=' + '+'.join(cnt_of(f'{get_column_letter(2+j)}{rr}') for j in range(4)) + '+1', GREEN)
put(tb, 'A21', 'Points so far', BOLD); put(tb, 'G21', '=SUM(G17:G20)')
put(tb, 'H21', '=IF(G21=B5+B6,"matches bidders + defenders","check")', GREY)
put(tb, 'H17', 'Who sat where for tricks 1–4 is not part of the trace; only the tiles and the running score are.', GREY)
put(tb, 'A23', 'Unseen by seat 2 (8 tiles)', BOLD)
for i, t in enumerate(data['unseen']): put(tb, f'{get_column_letter(2+i)}23', t, BLUE)
put(tb, 'A24', 'Check: 16 played + 3 in hand + 1 led + 8 unseen', GREY); put(tb, 'B24', '=COUNTA(B17:E20)+COUNTA(B10:D10)+COUNTA(B11)+COUNTA(B23:I23)')
put(tb, 'C24', '=IF(B24=28,"= 28 ✓","≠ 28")')

# ---------------------------------------------------------------- 4. Worlds
wo = wb.create_sheet('Worlds')
widths(wo, [8, 9, 9, 9, 9, 9, 9, 9, 9, 10, 50])
put(wo, 'A1', 'Eight guesses at the hidden deal', H2)
put(wo, 'A2', 'Each world hands the 8 unseen tiles to seats 0, 1 and 3 (3, 2 and 3 tiles — the sizes are known from the tiles already played). A world stays alive as long as nothing seen so far rules it out.', GREY)
header(wo, 4, ['World', 'Seat 0', '', '', 'Seat 1 (walt)', '', 'Seat 3', '', '', 'Check', 'Note'])
WORLD_ROW = {}
for w, deal in enumerate(data['deals']):
    rr = 5 + w; WORLD_ROW[w] = rr
    put(wo, f'A{rr}', f'W{w+1}', BOLD)
    cells = deal['0'] + deal['1'] + deal['3']
    for i, t in enumerate(cells): put(wo, f'{get_column_letter(2+i)}{rr}', t, BLUE)
    put(wo, f'J{rr}', f'=IF(AND(COUNTA(B{rr}:I{rr})=8,SUMPRODUCT(COUNTIF(B{rr}:I{rr},Table!$B$23:$I$23))=8),"all 8 once ✓","check")')
wo.merge_cells('B4:D4'); wo.merge_cells('E4:F4'); wo.merge_cells('G4:I4')
put(wo, 'K5', 'Seat 3 holds the 6-6 in five of eight worlds (W1, W2, W5, W6, W8).', GREY)
put(wo, 'K6', 'Seat 2 is seat 2 in every world: 6-3, 6-0, 5-5.', GREY)
put(wo, 'A14', 'How many worlds hold a given tile at a given seat:', BOLD)
header(wo, 15, ['Tile', 'Seat 0', 'Seat 1', 'Seat 3'])
for i, t in enumerate(data['unseen']):
    rr = 16 + i
    put(wo, f'A{rr}', t, BLUE)
    put(wo, f'B{rr}', f'=COUNTIF($B$5:$D$12,A{rr})')
    put(wo, f'C{rr}', f'=COUNTIF($E$5:$F$12,A{rr})')
    put(wo, f'D{rr}', f'=COUNTIF($G$5:$I$12,A{rr})')
wo.freeze_panes = 'B5'

# ---------------------------------------------------------------- 5. Trees
TREE_COLS = ['Node', 'Depth', 'Parent', 'Kind', 'Seat to play', 'Reached by', 'Trick so far',
             'Worlds alive', '# worlds', 'Bidders pts', 'Defenders pts', 'Value\n(worlds where bid is made)',
             'Leaf result', 'Kept by the plan node above?', 'How the value is computed']
ROW_OF = {}   # html id -> (sheet title, row)
TEAM = {0: 'defender', 1: 'bidder (walt)', 2: 'seat 2 — the mind', 3: 'bidder'}

def build_tree(title, prefix, raw, html_root, tile):
    ws = wb.create_sheet(title)
    widths(ws, [16, 6, 14, 9, 9, 10, 22, 24, 8, 9, 9, 13, 11, 14, 46])
    header(ws, 4, TREE_COLS)
    ws.row_dimensions[4].height = 42
    put(ws, 'A1', f'{title}: seat 2 plays the {tile} — every position, all the way down', H2)
    put(ws, 'A2', 'Rows are positions. A parent\'s Value is a formula over the rows directly under it (its children). Collapse the outline to watch the counts climb.', GREY)
    rows = []  # list of dicts
    def walk(node, nid, parent, tile, depth, html_id):
        rec = dict(id=nid, parent=parent, tile=tile, depth=depth, html=html_id, raw=node)
        rows.append(rec)
        for i, ch in enumerate(node.get('ch', [])):
            cid = f'{nid}.{i+1}'
            if ch.get('skip'):
                rows.append(dict(id=cid, parent=nid, tile=ch['tile'], depth=depth+1, html=f'{html_id}.{i}', raw=None))
            else:
                walk(ch['n'], cid, nid, ch['tile'], depth+1, f'{html_id}.{i}')
    walk(raw, prefix, 'root', prefix, 0, html_root)
    first = 5; last = first + len(rows) - 1
    V = f'$L${first}:$L${last}'; P = f'$C${first}:$C${last}'; A = f'$A${first}:$A${last}'
    for k, rec in enumerate(rows):
        r = first + k; raw = rec['raw']; rid = rec['id']
        ROW_OF[rec['html']] = (title, r)
        ind = '    ' * min(rec['depth'], 10)
        put(ws, f'A{r}', ind + rid, BOLD if rec['depth'] == 0 else BLACK)
        put(ws, f'B{r}', rec['depth'])
        put(ws, f'C{r}', rec['parent'])
        put(ws, f'F{r}', rec['tile'], BLUE)
        if raw is None:
            put(ws, f'D{r}', 'skipped', GREY)
            put(ws, f'O{r}', 'Never visited: a sibling already scored 0 and nothing beats zero for a defender taking the MIN.', GREY)
            ws.row_dimensions[r].outlineLevel = min(rec['depth'], 7)
            continue
        kind = 'plan' if raw['k'] == 'own' else raw['k']
        put(ws, f'D{r}', kind, fill=PLAN_FILL if kind == 'plan' else None)
        if kind != 'leaf': put(ws, f'E{r}', f"S{raw['seat']}", BLUE)
        put(ws, f'G{r}', ' → '.join(raw['trick']) if raw['trick'] else '(new trick)', BLUE)
        put(ws, f'H{r}', ', '.join(f'W{w+1}' for w in raw['w']), BLUE)
        put(ws, f'I{r}', f'=LEN(H{r})-LEN(SUBSTITUTE(H{r},"W",""))')
        put(ws, f'J{r}', raw['t1'], BLUE); put(ws, f'K{r}', raw['t0'], BLUE)
        if kind == 'dice':
            put(ws, f'O{r}', f"S{raw['seat']} rolls a die over its legal tiles in each world. SUM of the children: the roll split the worlds into piles.")
        elif kind == 'plan':
            put(ws, f'O{r}', 'Seat 2\'s own turn: no die. Try every legal tile in all alive worlds at once. MIN of the children (defending).')
        else:
            put(ws, f'L{r}', f'=IF(J{r}>=Table!$B$4,I{r},0)')
            put(ws, f'M{r}', f'=IF(J{r}>=Table!$B$4,"bid made",IF(K{r}>=Table!$B$7,"bid set",""))')
            put(ws, f'O{r}', 'Hand decided. Value = number of worlds here if the bidders reached the bid, else 0.')
        ws.row_dimensions[r].outlineLevel = min(rec['depth'], 7)
    # helper column P holds the untrimmed node id
    for k, rec in enumerate(rows):
        r = first + k
        put(ws, f'P{r}', rec['id'], GREY)
    ws.column_dimensions['P'].hidden = True
    def sub_end(k):
        d = rows[k]['depth']; j = k + 1
        while j < len(rows) and rows[j]['depth'] > d: j += 1
        return first + j - 1
    for k, rec in enumerate(rows):
        r = first + k
        if rec['raw'] is None: continue
        e = sub_end(k); SV = f'$L${r+1}:$L${e}'; SP = f'$C${r+1}:$C${e}'
        if rec['parent'] != 'root': put(ws, f'N{r}', f'=IF(INDEX($D${first}:$D${last},MATCH(C{r},$P${first}:$P${last},0))="plan",IF(L{r}=INDEX({V},MATCH(C{r},$P${first}:$P${last},0)),"kept","dropped"),"")')
        if rec['raw']['k'] == 'dice':
            put(ws, f'L{r}', f'=SUMIFS({SV},{SP},P{r})')
        elif rec['raw']['k'] == 'own':
            put(ws, f'L{r}', f'=IF(Table!$B$8="defending",_xlfn.MINIFS({SV},{SP},P{r}),_xlfn.MAXIFS({SV},{SP},P{r}))', fill=PLAN_FILL)
    # footer
    fr = last + 2
    put(ws, f'A{fr}', 'Positions visited', BOLD); put(ws, f'L{fr}', f'=COUNTIF($D${first}:$D${last},"dice")+COUNTIF($D${first}:$D${last},"plan")+COUNTIF($D${first}:$D${last},"leaf")')
    put(ws, f'A{fr+1}', 'Leaves where the bid was made', BOLD); put(ws, f'L{fr+1}', f'=COUNTIF($M${first}:$M${last},"bid made")')
    put(ws, f'A{fr+2}', 'Leaves where the bid was set', BOLD); put(ws, f'L{fr+2}', f'=COUNTIF($M${first}:$M${last},"bid set")')
    put(ws, f'A{fr+3}', f'Worlds where the bid is made if seat 2 plays the {tile}', BOLD); put(ws, f'L{fr+3}', f'=L{first}', BOLD)
    # conditional formats
    rng = f'A{first}:O{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{first}="bid made"'], fill=MADE_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$M{first}="bid set"'], fill=SET_FILL))
    ws.conditional_formatting.add(f'N{first}:N{last}', FormulaRule(formula=[f'$N{first}="dropped"'], font=Font(name=F, color='999999', italic=True)))
    ws.conditional_formatting.add(f'N{first}:N{last}', FormulaRule(formula=[f'$N{first}="kept"'], font=Font(name=F, color='7A3E97', bold=True)))
    ws.freeze_panes = 'B5'
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    for r in range(first, last + 1):
        ws[f'G{r}'].alignment = Alignment(shrink_to_fit=True)
    return ws, first, fr + 3

t63, T63_FIRST, T63_SUM = build_tree('Tree 6-3', 'A', data['res']['6-3']['tree'], 'r.0', '6-3')
t60, T60_FIRST, T60_SUM = build_tree('Tree 6-0', 'B', data['res']['6-0']['tree'], 'r.1', '6-0')

# ---------------------------------------------------------------- 6. Decision
dc = wb.create_sheet('Decision', 1)
widths(dc, [34, 14, 14, 60])
put(dc, 'A1', 'The decision: 6-3 or 6-0?', H1)
put(dc, 'A2', 'The root is a plan node (seat 2\'s own turn). It tried both legal tiles; each child is a whole tree on its own sheet.', GREY)
header(dc, 4, ['', 'Play the 6-3', 'Play the 6-0', 'Note'])
put(dc, 'A5', 'Worlds where the bid is made', BOLD)
put(dc, 'B5', f"='Tree 6-3'!L{T63_FIRST}", GREEN); put(dc, 'C5', f"='Tree 6-0'!L{T60_FIRST}", GREEN)
put(dc, 'D5', 'This pair of counts is the whole output of the search: the vector.', GREY)
put(dc, 'A6', 'Positions it took to find out', BOLD)
put(dc, 'B6', f"='Tree 6-3'!L{T63_SUM-3}", GREEN); put(dc, 'C6', f"='Tree 6-0'!L{T60_SUM-3}", GREEN)
put(dc, 'D6', '=B6+C6&" positions in total, for one guess at the table."', GREY)
put(dc, 'A8', 'Seat 2 is', BOLD); put(dc, 'B8', '=Table!B8', GREEN)
put(dc, 'A9', 'So the root takes the', BOLD); put(dc, 'B9', '=IF(B8="defending","MIN","MAX")')
put(dc, 'A10', 'Root value', BOLD); put(dc, 'B10', '=IF(B8="defending",MIN(B5:C5),MAX(B5:C5))', BOLD)
put(dc, 'A11', 'Seat 2 plays', BOLD)
put(dc, 'B11', '=IF(B5=C5,"6-0 (tie)",IF(B8="defending",IF(B5<C5,"6-3","6-0"),IF(B5>C5,"6-3","6-0")))', BOLD)
put(dc, 'D11', 'Tie rule: ties go to the lower tile index, the 6-0. Why that matters is part one, "One Mind, Eight Worlds".', GREY)
put(dc, 'A13', 'Change Table!B8 to "bidding" and the whole tree flips from MIN to MAX in every plan node — the same 151 rows answer the opposite question.', GREY)
put(dc, 'A15', 'One of 2,900', H2)
put(dc, 'A16', 'This tree was one mind\'s answer, in one of walt\'s sampled worlds. Before it plays one tile at trick 1, walt asks about 2,900 of them, and each is two functions and a count: roll for everyone else, try everything for yourself, count the worlds where the bid is made.', BLACK, al=WRAP)
dc.merge_cells('A16:D16'); dc.row_dimensions[16].height = 48
dc.sheet_view.showGridLines = False

# ---------------------------------------------------------------- 7. Walkthrough
wk = wb.create_sheet('Walkthrough', 2)
widths(wk, [6, 22, 24, 20, 90])
put(wk, 'A1', 'Walkthrough: the five acts', H1)
put(wk, 'A2', 'Read the caption, click the link, look at the row. Captions follow the animated explainer scene for scene.', GREY)
header(wk, 4, ['Act', 'Scene', 'Look at', 'Sheet', 'Caption'])
ACTS = {1: 'What a node is', 2: 'Two kinds of node', 3: 'Building the 6-3 tree', 4: 'The other candidate', 5: 'Zoom out'}
# scene -> html node id (or 'decision'/'worlds'/'table')
look = {
 'A node': 'r', 'The bar': 'table', 'What alive means': 'worlds', 'The question': 'r',
 'Three shapes': 'readme', 'Dice node': 'r.1.1', 'Sum': 'r.1.1.0', 'Plan node': 'r.0.0.1.0.0.0', 'Min': 'r.0.0.1.0.0.0',
 'Leaf': 'r.0.0.0.0.0.0.1',
 'Try the 6-3': 'r.0', 'Seat 3 rolls': 'r.0', 'Seat 0 rolls': 'r.0.0', 'Follow W2': 'r.0.0.1', 'No dice': 'r.0.0.1.0.0.0',
 'Trump in': 'r.0.0.1.0.0.0.0', 'Throw off': 'r.0.0.1.0.0.0.1', 'Take the min': 'r.0.0.1.0.0.0', 'Speeding up': 'r.0.0.0',
 'Early exit': 'r.0.0.2.0.0', 'Three of five': 'r.0.0', 'The small groups': 'r.0.1', 'Five of eight': 'r.0',
 'The other candidate': 'r.1', 'Side by side': 'decision', 'A tie': 'decision', 'All of it': 'decision', 'One of 2,900': 'decision',
}
r = 5; cur_act = None
for act, title, cap in scenes:
    if act != cur_act:
        cur_act = act
        put(wk, f'A{r}', f'Act {act}', BOLD); put(wk, f'B{r}', ACTS[act], BOLD)
        for col in 'ABCDE': wk[f'{col}{r}'].fill = HEAD_FILL
        r += 1
    put(wk, f'A{r}', act); put(wk, f'B{r}', title, BOLD)
    key = look[title]
    if key in ROW_OF:
        sh, rr = ROW_OF[key]
        nid = (t63 if sh == 'Tree 6-3' else t60)[f'P{rr}'].value
        q = "'"
        put(wk, f'C{r}', '=HYPERLINK("#' + q + sh + q + f'!A{rr}","{nid}  (row {rr})")', GREEN); put(wk, f'D{r}', sh)
    elif key == 'r':
        put(wk, f'C{r}', '=HYPERLINK("#Decision!A5","the root")', GREEN); put(wk, f'D{r}', 'Decision')
    elif key == 'decision':
        put(wk, f'C{r}', '=HYPERLINK("#Decision!B5","B5:C5")', GREEN); put(wk, f'D{r}', 'Decision')
    elif key == 'worlds':
        put(wk, f'C{r}', '=HYPERLINK("#Worlds!A5","W1–W8")', GREEN); put(wk, f'D{r}', 'Worlds')
    elif key == 'table':
        put(wk, f'C{r}', '=HYPERLINK("#Table!B5","the score")', GREEN); put(wk, f'D{r}', 'Table')
    elif key == 'readme':
        put(wk, f'C{r}', '=HYPERLINK("#' + "'Read me'" + '!B12","the three kinds")', GREEN); put(wk, f'D{r}', 'Read me')
    put(wk, f'E{r}', cap, al=WRAP)
    wk.row_dimensions[r].height = 15 * max(1, (len(cap) // 95) + 1)
    r += 1
wk.freeze_panes = 'A5'

# ---------------------------------------------------------------- 8. Two rules (the engine primitives)
tr = wb.create_sheet('Two rules')
widths(tr, [30, 12, 12, 12, 12, 12, 12, 12, 48])
put(tr, 'A1', 'The two rules Walt needs from the game', H1)
put(tr, 'A2', 'Everything else in the tree is counting. Change the yellow cells and the black cells recompute.', GREY)
tiles_rng = f'Tiles!$A$5:$A${TILE_LAST}'
def lk(col, ref): return f'INDEX(Tiles!${col}$5:${col}${TILE_LAST},MATCH({ref},{tiles_rng},0))'
# Rule 1: who wins a trick
put(tr, 'A4', 'Rule 1 — who wins a trick, and what it is worth', H2)
header(tr, 5, ['', 'Lead', '2nd', '3rd', '4th', '', '', '', 'Note'])
put(tr, 'A6', 'Tiles played (in order)', BOLD)
for j, t in enumerate(['6-4', '6-3', '6-6', '5-2']): put(tr, f'{get_column_letter(2+j)}6', t, BLUE, fill=INPUT_FILL)
put(tr, 'A7', 'Suit when led'); put(tr, 'A8', 'Rank when led'); put(tr, 'A9', 'Count points'); put(tr, 'A10', 'Strength in this trick')
for j in range(4):
    c = get_column_letter(2 + j); ref = f'{c}$6'
    put(tr, f'{c}7', f'={lk("D", ref)}', GREEN)
    put(tr, f'{c}8', f'={lk("E", ref)}', GREEN)
    put(tr, f'{c}9', f'={lk("F", ref)}', GREEN)
    put(tr, f'{c}10', f'=IF(OR({lk("B", ref)}=5,{lk("C", ref)}=5),100+IF({lk("B", ref)}={lk("C", ref)},7,{lk("B", ref)}+{lk("C", ref)}-5),IF($B$7="fives (trump)",0,IF(OR({lk("B", ref)}=$B$14,{lk("C", ref)}=$B$14),IF({lk("B", ref)}={lk("C", ref)},7,{lk("B", ref)}+{lk("C", ref)}-$B$14),0)))')
put(tr, 'I10', 'Trump beats everything. Otherwise a tile counts only if it contains the led pip, and then its other end is its rank (double highest). Off-suit = 0.', GREY)
put(tr, 'A11', 'Winning tile', BOLD); put(tr, 'B11', '=INDEX(B6:E6,MATCH(MAX(B10:E10),B10:E10,0))', BOLD)
put(tr, 'A12', 'Trick is worth', BOLD); put(tr, 'B12', '=SUM(B9:E9)+1', BOLD); put(tr, 'I12', 'Count points on the tiles, plus 1 for the trick itself.', GREY)
put(tr, 'A13', 'Winner\'s position', BOLD); put(tr, 'B13', '=MATCH(MAX(B10:E10),B10:E10,0)'); put(tr, 'I13', '1 = the leader, 2 = next player, and so on.', GREY)
put(tr, 'A14', 'Led pip (if not trump)', BOLD); put(tr, 'B14', f'={lk("B","B6")}', GREEN)
# Rule 2: which tiles may be played
put(tr, 'A16', 'Rule 2 — which tiles a hand may play', H2)
put(tr, 'A17', 'Tile led this trick', BOLD); put(tr, 'B17', '6-4', BLUE, fill=INPUT_FILL)
put(tr, 'C17', 'leave blank if this hand is leading', GREY)
put(tr, 'A18', 'Led suit'); put(tr, 'B18', f'=IF(B17="","(leading)",{lk("D","B17")})', GREEN)
header(tr, 20, ['', 'Tile 1', 'Tile 2', 'Tile 3', 'Tile 4', 'Tile 5', 'Tile 6', 'Tile 7', 'Note'])
put(tr, 'A21', 'Hand', BOLD)
for j, t in enumerate(['6-3', '6-0', '5-5', '', '', '', '']): put(tr, f'{get_column_letter(2+j)}21', t if t else None, BLUE, fill=INPUT_FILL)
put(tr, 'A22', 'Suit')
put(tr, 'A23', 'Follows the led suit?')
put(tr, 'A24', 'May be played')
for j in range(7):
    c = get_column_letter(2 + j)
    put(tr, f'{c}22', f'=IF({c}21="","",{lk("D", c + "21")})', GREEN)
    put(tr, f'{c}23', f'=IF({c}21="","",IF($B$17="","n/a",IF($B$18="fives (trump)",IF({c}22="fives (trump)","yes","no"),IF({c}22="fives (trump)","no",IF(OR({lk("B", c + "21")}={lk("B","B17")},{lk("C", c + "21")}={lk("B","B17")}),"yes","no")))))')
    put(tr, f'{c}24', f'=IF({c}21="","",IF(OR($B$17="",$B$25=0,{c}23="yes"),"✓",""))', BOLD)
put(tr, 'A25', 'Tiles in hand that follow suit'); put(tr, 'B25', '=COUNTIF(B23:H23,"yes")')
put(tr, 'I24', 'A tile follows if it contains the led pip (either end) and is not trump; if trump was led only trumps follow. If any tile follows you must play one; if none do, anything goes. Leading: anything goes.', GREY)
put(tr, 'A28', 'Where a full Excel Walt goes from here', H2)
notes = [
 'A dice node needs Rule 2 (legal tiles for that seat in that world) and a die.',
 'A plan node needs Rule 2 for seat 2 and a MIN (or MAX).',
 'A leaf needs Rule 1 four times per trick, to keep score until someone crosses 30 or 13.',
 'The tree sheets here are a trace: the recursion was done elsewhere and written down row by row so each Value is a formula over its children.',
 'To make Excel play by itself, something has to generate the rows: a recursive LAMBDA (Excel 365) or a short VBA routine. Both are a page of code. The rules above are already in formulas.',
]
for i, n in enumerate(notes):
    put(tr, f'A{29+i}', f'{i+1}.  {n}', al=WRAP); tr.merge_cells(f'A{29+i}:I{29+i}')

exec(open('flat.py').read())
# sheet order
order = ['Read me', 'Decision', 'Walkthrough', 'Table', 'Worlds', 'Tree 6-3', 'Tree 6-0', 'Two rules', 'Tiles', 'Tape'] + LINE_SHEETS + ['Flat Walt', 'Swap votes']
wb._sheets = [wb[n] for n in order]
wb.save('/mnt/user-data/outputs/Walt in Excel - All the Way Down.xlsx')
print('saved', ROW_OF.get('r.0.0.1.0.0.0'))
